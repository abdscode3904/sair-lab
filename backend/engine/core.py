from __future__ import annotations

from pathlib import Path
from typing import Optional, Iterable
import math
import re
import shutil
import zipfile

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.formatting.rule import ColorScaleRule


# ============================================================
# SAIR LAB EXCEL AUTOMATION ENGINE
# Day 1 - Core Processing Engine
# ============================================================


SUPPORTED_INPUTS = {
    ".xlsx",
    ".xlsm",
    ".csv",
}

SUPPORTED_OUTPUTS = {
    ".xlsx",
    ".csv",
}


class SairLabError(Exception):
    """Base exception for Sair Lab processing errors."""


class FileValidationError(SairLabError):
    """Raised when an input file is invalid."""


class ProcessingError(SairLabError):
    """Raised when an automation operation fails."""


# ============================================================
# FILE VALIDATION
# ============================================================

def validate_input_file(file_path: str | Path) -> Path:
    path = Path(file_path)

    if not path.exists():
        raise FileValidationError(f"File does not exist: {path}")

    if not path.is_file():
        raise FileValidationError(f"Not a file: {path}")

    if path.suffix.lower() not in SUPPORTED_INPUTS:
        raise FileValidationError(
            f"Unsupported file type: {path.suffix}. "
            f"Supported: {sorted(SUPPORTED_INPUTS)}"
        )

    if path.stat().st_size == 0:
        raise FileValidationError("Input file is empty.")

    return path


def validate_output_path(output_path: str | Path) -> Path:
    path = Path(output_path)

    if path.suffix.lower() not in SUPPORTED_OUTPUTS:
        raise FileValidationError(
            f"Unsupported output type: {path.suffix}"
        )

    path.parent.mkdir(parents=True, exist_ok=True)

    return path


# ============================================================
# LOADERS
# ============================================================

def load_csv(file_path: str | Path) -> pd.DataFrame:
    path = validate_input_file(file_path)

    try:
        return pd.read_csv(path)
    except Exception as exc:
        raise ProcessingError(
            f"Could not read CSV file: {exc}"
        ) from exc


def load_excel(
    file_path: str | Path,
    sheet_name: str | int = 0,
) -> pd.DataFrame:
    path = validate_input_file(file_path)

    try:
        return pd.read_excel(path, sheet_name=sheet_name)
    except Exception as exc:
        raise ProcessingError(
            f"Could not read Excel file: {exc}"
        ) from exc


def load_table(
    file_path: str | Path,
    sheet_name: str | int = 0,
) -> pd.DataFrame:
    path = validate_input_file(file_path)

    if path.suffix.lower() == ".csv":
        return load_csv(path)

    return load_excel(path, sheet_name)


# ============================================================
# BASIC DATA CLEANING
# ============================================================

def remove_empty_rows(df: pd.DataFrame) -> pd.DataFrame:
    return df.dropna(how="all").reset_index(drop=True)


def remove_empty_columns(df: pd.DataFrame) -> pd.DataFrame:
    return df.dropna(axis=1, how="all")


def trim_whitespace(df: pd.DataFrame) -> pd.DataFrame:
    result = df.copy()

    for column in result.columns:
        if (
            pd.api.types.is_object_dtype(result[column])
            or pd.api.types.is_string_dtype(result[column])
        ):
            result[column] = result[column].apply(
                lambda value: value.strip()
                if isinstance(value, str)
                else value
            )

    return result


def normalize_text(
    df: pd.DataFrame,
    columns: Optional[Iterable[str]] = None,
    lowercase: bool = False,
    uppercase: bool = False,
) -> pd.DataFrame:

    result = df.copy()

    selected = list(columns) if columns else list(result.columns)

    for column in selected:
        if column not in result.columns:
            continue

        if (
            pd.api.types.is_object_dtype(result[column])
            or pd.api.types.is_string_dtype(result[column])
        ):
            series = result[column].astype("string").str.strip()

            if lowercase:
                series = series.str.lower()

            if uppercase:
                series = series.str.upper()

            result[column] = series

    return result


def remove_duplicates(
    df: pd.DataFrame,
    subset: Optional[list[str]] = None,
) -> pd.DataFrame:

    if subset:
        missing = [column for column in subset if column not in df.columns]

        if missing:
            raise ProcessingError(
                f"Duplicate-check columns not found: {missing}"
            )

    return df.drop_duplicates(
        subset=subset,
        keep="first",
    ).reset_index(drop=True)


def fill_missing_values(
    df: pd.DataFrame,
    value: object = "",
) -> pd.DataFrame:

    return df.fillna(value)


def clean_dataframe(df: pd.DataFrame) -> pd.DataFrame:
    result = df.copy()

    result = remove_empty_rows(result)
    result = remove_empty_columns(result)
    result = trim_whitespace(result)
    result = remove_duplicates(result)

    return result.reset_index(drop=True)


# ============================================================
# COLUMN OPERATIONS
# ============================================================

def rename_columns(
    df: pd.DataFrame,
    mapping: dict[str, str],
) -> pd.DataFrame:

    result = df.copy()

    missing = [
        old_name
        for old_name in mapping
        if old_name not in result.columns
    ]

    if missing:
        raise ProcessingError(
            f"Columns not found: {missing}"
        )

    return result.rename(columns=mapping)


def select_columns(
    df: pd.DataFrame,
    columns: list[str],
) -> pd.DataFrame:

    missing = [
        column for column in columns
        if column not in df.columns
    ]

    if missing:
        raise ProcessingError(
            f"Columns not found: {missing}"
        )

    return df[columns].copy()


def reorder_columns(
    df: pd.DataFrame,
    columns: list[str],
) -> pd.DataFrame:

    existing = list(df.columns)

    requested = [column for column in columns if column in existing]

    remaining = [
        column for column in existing
        if column not in requested
    ]

    return df[requested + remaining].copy()


def delete_columns(
    df: pd.DataFrame,
    columns: list[str],
) -> pd.DataFrame:

    return df.drop(
        columns=columns,
        errors="ignore",
    )


def split_column(
    df: pd.DataFrame,
    column: str,
    separator: str,
    new_columns: Optional[list[str]] = None,
) -> pd.DataFrame:

    if column not in df.columns:
        raise ProcessingError(
            f"Column not found: {column}"
        )

    result = df.copy()

    split_data = result[column].astype("string").str.split(
        separator,
        expand=True,
    )

    if new_columns:
        if len(new_columns) != split_data.shape[1]:
            raise ProcessingError(
                "Number of new column names does not match "
                "the number of split columns."
            )

        split_data.columns = new_columns
    else:
        split_data.columns = [
            f"{column}_{index + 1}"
            for index in range(split_data.shape[1])
        ]

    result = result.drop(columns=[column])

    return pd.concat(
        [result, split_data],
        axis=1,
    )


def merge_columns(
    df: pd.DataFrame,
    columns: list[str],
    new_column: str,
    separator: str = " ",
) -> pd.DataFrame:

    missing = [
        column for column in columns
        if column not in df.columns
    ]

    if missing:
        raise ProcessingError(
            f"Columns not found: {missing}"
        )

    result = df.copy()

    result[new_column] = (
        result[columns]
        .fillna("")
        .astype(str)
        .agg(separator.join, axis=1)
        .str.strip()
    )

    return result


# ============================================================
# SORTING / FILTERING
# ============================================================

def sort_dataframe(
    df: pd.DataFrame,
    column: str,
    ascending: bool = True,
) -> pd.DataFrame:

    if column not in df.columns:
        raise ProcessingError(
            f"Column not found: {column}"
        )

    return df.sort_values(
        by=column,
        ascending=ascending,
        kind="stable",
    ).reset_index(drop=True)


def filter_dataframe(
    df: pd.DataFrame,
    column: str,
    value: object,
) -> pd.DataFrame:

    if column not in df.columns:
        raise ProcessingError(
            f"Column not found: {column}"
        )

    return df[df[column] == value].reset_index(drop=True)


def filter_contains(
    df: pd.DataFrame,
    column: str,
    text: str,
) -> pd.DataFrame:

    if column not in df.columns:
        raise ProcessingError(
            f"Column not found: {column}"
        )

    mask = (
        df[column]
        .astype("string")
        .str.contains(
            re.escape(text),
            case=False,
            na=False,
        )
    )

    return df[mask].reset_index(drop=True)


# ============================================================
# FIND / REPLACE
# ============================================================

def find_replace(
    df: pd.DataFrame,
    old_value: object,
    new_value: object,
) -> pd.DataFrame:

    return df.replace(
        old_value,
        new_value,
    )


# ============================================================
# DATA TYPES
# ============================================================

def convert_numeric(
    df: pd.DataFrame,
    columns: list[str],
) -> pd.DataFrame:

    result = df.copy()

    for column in columns:
        if column not in result.columns:
            raise ProcessingError(
                f"Column not found: {column}"
            )

        result[column] = pd.to_numeric(
            result[column],
            errors="coerce",
        )

    return result


def convert_dates(
    df: pd.DataFrame,
    columns: list[str],
) -> pd.DataFrame:

    result = df.copy()

    for column in columns:
        if column not in result.columns:
            raise ProcessingError(
                f"Column not found: {column}"
            )

        result[column] = pd.to_datetime(
            result[column],
            errors="coerce",
        )

    return result


# ============================================================
# EXCEL EXPORT
# ============================================================

def export_excel(
    df: pd.DataFrame,
    output_path: str | Path,
    sheet_name: str = "Data",
    auto_width: bool = True,
) -> Path:

    output = validate_output_path(output_path)

    try:
        with pd.ExcelWriter(
            output,
            engine="xlsxwriter",
            datetime_format="yyyy-mm-dd",
        ) as writer:

            df.to_excel(
                writer,
                index=False,
                sheet_name=sheet_name,
            )

            workbook = writer.book
            worksheet = writer.sheets[sheet_name]

            header_format = workbook.add_format(
                {
                    "bold": True,
                    "border": 1,
                    "align": "center",
                    "valign": "vcenter",
                }
            )

            cell_format = workbook.add_format(
                {
                    "border": 1,
                    "valign": "top",
                }
            )

            for column_number, column_name in enumerate(df.columns):
                worksheet.write(
                    0,
                    column_number,
                    column_name,
                    header_format,
                )

                worksheet.set_column(
                    column_number,
                    column_number,
                    min(
                        max(
                            len(str(column_name)) + 2,
                            12,
                        ),
                        45,
                    ),
                    cell_format,
                )

            worksheet.freeze_panes(1, 0)
            worksheet.autofilter(
                0,
                0,
                max(len(df), 1),
                max(len(df.columns) - 1, 0),
            )

    except Exception as exc:
        raise ProcessingError(
            f"Could not export Excel file: {exc}"
        ) from exc

    return output


def export_csv(
    df: pd.DataFrame,
    output_path: str | Path,
) -> Path:

    output = validate_output_path(output_path)

    try:
        df.to_csv(
            output,
            index=False,
        )
    except Exception as exc:
        raise ProcessingError(
            f"Could not export CSV file: {exc}"
        ) from exc

    return output


# ============================================================
# EXCEL FORMATTING
# ============================================================

def format_excel(
    input_path: str | Path,
    output_path: str | Path,
    freeze_header: bool = True,
    auto_filter: bool = True,
    auto_width: bool = True,
) -> Path:

    input_file = validate_input_file(input_path)
    output_file = validate_output_path(output_path)

    try:
        shutil.copy2(
            input_file,
            output_file,
        )

        workbook = load_workbook(
            output_file
        )

        thin = Side(
            style="thin"
        )

        border = Border(
            left=thin,
            right=thin,
            top=thin,
            bottom=thin,
        )

        header_fill = PatternFill(
            fill_type="solid",
            fgColor="1F4E78",
        )

        header_font = Font(
            bold=True,
            color="FFFFFF",
        )

        for worksheet in workbook.worksheets:

            if worksheet.max_row == 0:
                continue

            for cell in worksheet[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.border = border
                cell.alignment = Alignment(
                    horizontal="center",
                    vertical="center",
                )

            for row in worksheet.iter_rows(
                min_row=2,
                max_row=worksheet.max_row,
                min_col=1,
                max_col=worksheet.max_column,
            ):
                for cell in row:
                    cell.border = border

            if freeze_header:
                worksheet.freeze_panes = "A2"

            if auto_filter:
                worksheet.auto_filter.ref = (
                    worksheet.dimensions
                )

            if auto_width:
                for column_cells in worksheet.columns:

                    max_length = 0

                    for cell in column_cells:
                        value = cell.value

                        if value is None:
                            continue

                        max_length = max(
                            max_length,
                            len(str(value)),
                        )

                    width = min(
                        max(max_length + 2, 12),
                        50,
                    )

                    column_letter = (
                        column_cells[0].column_letter
                    )

                    worksheet.column_dimensions[
                        column_letter
                    ].width = width

        workbook.save(output_file)

    except Exception as exc:
        raise ProcessingError(
            f"Could not format Excel file: {exc}"
        ) from exc

    return output_file


# ============================================================
# SUMMARY / REPORT
# ============================================================

def create_summary(
    df: pd.DataFrame,
) -> pd.DataFrame:

    summary_rows = []

    for column in df.columns:

        series = df[column]

        row = {
            "Column": str(column),
            "Data Type": str(series.dtype),
            "Non-Empty": int(series.notna().sum()),
            "Missing": int(series.isna().sum()),
            "Unique Values": int(series.nunique(dropna=True)),
        }

        if pd.api.types.is_numeric_dtype(series):

            row["Minimum"] = (
                series.min()
                if not series.dropna().empty
                else None
            )

            row["Maximum"] = (
                series.max()
                if not series.dropna().empty
                else None
            )

            row["Average"] = (
                series.mean()
                if not series.dropna().empty
                else None
            )

        else:
            row["Minimum"] = None
            row["Maximum"] = None
            row["Average"] = None

        summary_rows.append(row)

    return pd.DataFrame(summary_rows)


def create_report(
    df: pd.DataFrame,
    output_path: str | Path,
    title: str = "Sair Lab Automated Report",
) -> Path:

    output = validate_output_path(output_path)

    summary = create_summary(df)

    try:
        with pd.ExcelWriter(
            output,
            engine="xlsxwriter",
        ) as writer:

            workbook = writer.book

            title_format = workbook.add_format(
                {
                    "bold": True,
                    "font_size": 18,
                }
            )

            section_format = workbook.add_format(
                {
                    "bold": True,
                    "font_size": 13,
                }
            )

            number_format = workbook.add_format(
                {
                    "num_format": "#,##0.00",
                }
            )

            worksheet = workbook.add_worksheet(
                "Summary"
            )

            worksheet.write(
                "A1",
                title,
                title_format,
            )

            worksheet.write(
                "A3",
                "Dataset Overview",
                section_format,
            )

            worksheet.write(
                "A4",
                "Rows",
            )

            worksheet.write(
                "B4",
                len(df),
            )

            worksheet.write(
                "A5",
                "Columns",
            )

            worksheet.write(
                "B5",
                len(df.columns),
            )

            worksheet.write(
                "A6",
                "Duplicate Rows",
            )

            worksheet.write(
                "B6",
                int(df.duplicated().sum()),
            )

            worksheet.write(
                "A8",
                "Column Analysis",
                section_format,
            )

            summary.to_excel(
                writer,
                index=False,
                sheet_name="Summary",
                startrow=8,
            )

            data_sheet = workbook.add_worksheet(
                "Data"
            )

            for column_index, column in enumerate(
                df.columns
            ):
                data_sheet.write(
                    0,
                    column_index,
                    column,
                )

            for row_index, row in enumerate(
                df.itertuples(index=False),
                start=1,
            ):
                for column_index, value in enumerate(row):

                    if pd.isna(value):
                        value = ""

                    data_sheet.write(
                        row_index,
                        column_index,
                        value,
                    )

            data_sheet.freeze_panes(1, 0)

            for index, column in enumerate(df.columns):
                width = min(
                    max(
                        len(str(column)) + 2,
                        12,
                    ),
                    40,
                )

                data_sheet.set_column(
                    index,
                    index,
                    width,
                )

    except Exception as exc:
        raise ProcessingError(
            f"Could not create report: {exc}"
        ) from exc

    return output


# ============================================================
# FILE MERGING
# ============================================================

def merge_csv_files(
    input_files: list[str | Path],
    output_path: str | Path,
) -> Path:

    if not input_files:
        raise ProcessingError(
            "No input files supplied."
        )

    frames = []

    for file_path in input_files:
        frames.append(
            load_csv(file_path)
        )

    merged = pd.concat(
        frames,
        ignore_index=True,
    )

    return export_csv(
        merged,
        output_path,
    )


def merge_excel_files(
    input_files: list[str | Path],
    output_path: str | Path,
) -> Path:

    if not input_files:
        raise ProcessingError(
            "No input files supplied."
        )

    frames = []

    for file_path in input_files:
        frames.append(
            load_excel(file_path)
        )

    merged = pd.concat(
        frames,
        ignore_index=True,
    )

    return export_excel(
        merged,
        output_path,
    )


# ============================================================
# CSV ↔ EXCEL CONVERSION
# ============================================================

def csv_to_excel(
    input_path: str | Path,
    output_path: str | Path,
) -> Path:

    df = load_csv(input_path)

    return export_excel(
        df,
        output_path,
    )


def excel_to_csv(
    input_path: str | Path,
    output_path: str | Path,
) -> Path:

    df = load_excel(input_path)

    return export_csv(
        df,
        output_path,
    )


# ============================================================
# DATASET INFORMATION
# ============================================================

def dataset_info(
    file_path: str | Path,
) -> dict:

    df = load_table(file_path)

    numeric_columns = [
        column
        for column in df.columns
        if pd.api.types.is_numeric_dtype(df[column])
    ]

    return {
        "rows": int(len(df)),
        "columns": int(len(df.columns)),
        "column_names": [
            str(column)
            for column in df.columns
        ],
        "duplicate_rows": int(
            df.duplicated().sum()
        ),
        "missing_cells": int(
            df.isna().sum().sum()
        ),
        "numeric_columns": [
            str(column)
            for column in numeric_columns
        ],
        "file_size_bytes": Path(file_path).stat().st_size,
    }


# ============================================================
# COMPLETE CLEANING WORKFLOW
# ============================================================

def automatic_clean(
    input_path: str | Path,
    output_path: str | Path,
) -> dict:

    original = load_table(input_path)

    original_rows = len(original)
    original_columns = len(original.columns)
    original_duplicates = int(
        original.duplicated().sum()
    )

    cleaned = clean_dataframe(
        original
    )

    output = export_excel(
        cleaned,
        output_path,
    )

    return {
        "input_rows": original_rows,
        "output_rows": len(cleaned),
        "input_columns": original_columns,
        "output_columns": len(cleaned.columns),
        "duplicates_removed": (
            original_duplicates
            - int(cleaned.duplicated().sum())
        ),
        "output": str(output),
    }
